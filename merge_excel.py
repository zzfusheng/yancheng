"""
合并水电费明细(1).xlsx 与 水电新.xlsx

逻辑：以 B列ID 做关联键，对每个ID进行交叉连接（笛卡尔积）。
      将水电新的 H,I,J,K 列（水电类型,单价,数量,金额）追加到水电费明细(1)对应行后面。
"""
import openpyxl
from collections import defaultdict

SRC1 = "水电费明细(1).xlsx"
SRC2 = "水电新.xlsx"
OUT  = "水电费明细_合并结果.xlsx"

print("读取文件中...")
wb1 = openpyxl.load_workbook(SRC1, read_only=True, data_only=True)
ws1 = wb1.active

wb2 = openpyxl.load_workbook(SRC2, read_only=True, data_only=True)
ws2 = wb2.active

# --- 读取水电新，按 ID（B列）分组 ---
# 只保留 H,I,J,K 列（索引 7,8,9,10，0-based）
sd_map = defaultdict(list)   # {id: [(水电类型, 单价, 数量, 金额), ...]}

rows2 = list(ws2.iter_rows(min_row=2, values_only=True))
for row in rows2:
    rid = row[1]  # B列 = ID
    if rid is None:
        continue
    sd_map[rid].append((row[7], row[8], row[9], row[10]))  # H,I,J,K

wb2.close()

print(f"水电新 共 {len(rows2)} 行数据，唯一ID {len(sd_map)} 个")

# --- 构建输出 ---
wb_out = openpyxl.Workbook(write_only=True)
ws_out = wb_out.create_sheet("Sheet1")

# 表头：水电费明细(1) 的 18 列 + 4 列
header1 = list(ws1.iter_rows(min_row=1, max_row=1, values_only=True))[0]
header_out = list(header1) + ["水电类型", "单价", "数量", "金额"]
ws_out.append(header_out)

matched = 0
unmatched = 0
total_out = 0

for row in ws1.iter_rows(min_row=2, values_only=True):
    rid = row[1]  # B列 = 单据id
    sd_rows = sd_map.get(rid)
    if sd_rows:
        for sd in sd_rows:
            ws_out.append(list(row) + list(sd))
            total_out += 1
        matched += 1
    else:
        # 无匹配：保留原行，4列留空
        ws_out.append(list(row) + [None, None, None, None])
        unmatched += 1
        total_out += 1

wb1.close()

print(f"写出文件 {OUT} ...")
wb_out.save(OUT)

print(f"完成！")
print(f"  水电费明细(1) 有效数据行数：{matched + unmatched}")
print(f"  其中匹配到水电新的行数：{matched}，未匹配：{unmatched}")
print(f"  输出总行数（不含表头）：{total_out}")
print(f"  输出文件：{OUT}")
